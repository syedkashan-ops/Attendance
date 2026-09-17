import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
import pandas as pd


def dataframe_to_excel_bytes(df, report_date):
    if df is None:
        df = pd.DataFrame()

    output = io.BytesIO()

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        if df.empty:
            columns = [
                "Visit ID",
                "Date",
                "Employee Code",
                "Employee Name",
                "Outlet Code",
                "Outlet Name",
                "IN Date/Time",
                "IN Latitude",
                "IN Longitude",
                "IN GPS Accuracy",
                "OUT Date/Time",
                "OUT Latitude",
                "OUT Longitude",
                "OUT GPS Accuracy",
                "IN-OUT GPS Distance (m)",
                "GPS Status",
                "Time Spent",
                "Time Spent Hours",
                "Status",
            ]
            pd.DataFrame(columns=columns).to_excel(
                writer,
                index=False,
                sheet_name="Visits",
            )
        else:
            df.to_excel(
                writer,
                index=False,
                sheet_name="Visits",
            )

        wb = writer.book
        ws = writer.sheets["Visits"]
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center")

        for col_cells in ws.columns:
            max_len = 0
            for cell in col_cells:
                value = "" if cell.value is None else str(cell.value)
                max_len = max(max_len, len(value))
            width = min(max(max_len + 2, 12), 35)
            ws.column_dimensions[get_column_letter(
                col_cells[0].column
            )].width = width

    return output.getvalue()


def management_excel_bytes(day_df, summary_df=None, employee_df=None, outlet_df=None, gps_df=None, daily_df=None):
    """Create a multi-sheet management workbook without changing the live Visits schema."""
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        sheets = [
            (day_df if day_df is not None else pd.DataFrame(), "Visits"),
            (summary_df if summary_df is not None else pd.DataFrame(), "Summary"),
            (employee_df if employee_df is not None else pd.DataFrame(), "Employee Analysis"),
            (outlet_df if outlet_df is not None else pd.DataFrame(), "Outlet Analysis"),
            (gps_df if gps_df is not None else pd.DataFrame(), "GPS Analysis"),
            (daily_df if daily_df is not None else pd.DataFrame(), "Daily Trend"),
        ]
        for frame, name in sheets:
            frame.to_excel(writer, index=False, sheet_name=name)
        wb = writer.book
        for ws in wb.worksheets:
            ws.freeze_panes = "A2"
            if ws.max_row >= 1 and ws.max_column >= 1:
                ws.auto_filter.ref = ws.dimensions
                for cell in ws[1]:
                    cell.font = Font(bold=True)
                    cell.alignment = Alignment(horizontal="center")
                for col_cells in ws.columns:
                    max_len = max((len(str(c.value)) if c.value is not None else 0) for c in col_cells)
                    ws.column_dimensions[get_column_letter(col_cells[0].column)].width = min(max(max_len + 2, 12), 35)
    return output.getvalue()
